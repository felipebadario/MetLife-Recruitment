"""Gera planilhas FICTÍCIAS no mesmo modelo das planilhas de turma (padrão "Base_selecao").

Serve para testes automatizados e para a base de demonstração. Nenhum dado aqui
corresponde a pessoas ou agências reais: nomes são combinações aleatórias de
prenomes e sobrenomes comuns, as agências têm nomes inventados e os CPFs são números
sorteados com dígito verificador válido, sem relação com ninguém.
"""

from __future__ import annotations

import io
import random
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.text import normalize_key

HEADER = [
    "Turma",
    "CPF Candidato",
    "Candidato",
    "R.N",
    "Nome Agência",
    "Agência",
    "Regional",
    "Executivo",
    "Suspenso",
    "Provas",
    "Status VGCF",
    "Status ENS",
    "Status Final",
    "Data Avaliação",
    "Data Cadastro",
]

DEMO_AGENCIES = [
    ("AURORA", "SUD"), ("BÚSSOLA", "SUD"), ("CEDRO", "SUL"), ("DUNAS", "CE-NO-NE"),
    ("ESTEIO", "SUD"), ("FAROL", "CE-NO-NE"), ("GALÁXIA", "SUD"), ("HORIZONTE", "SUL"),
    ("ÍCONE", "SUD"), ("JATOBÁ", "CE-NO-NE"), ("LUMEN", "SUD"), ("MARÉ", "CE-NO-NE"),
    ("NÚCLEO", "SUL"), ("ÓRBITA", "SUD"), ("PRISMA", "SUD"), ("QUARTZO", "SUL"),
    ("RUMO", "CE-NO-NE"), ("SEIVA", "SUD"), ("TERRA", "SUL"), ("URBANA", "SUD"),
    ("VÉRTICE", "CE-NO-NE"), ("XISTO", "SUD"), ("ZÊNITE", "SUL"),
]

FIRST_NAMES = [
    "Ana", "Bruno", "Carla", "Daniel", "Eduarda", "Felipe", "Gabriela", "Henrique", "Isabela",
    "João", "Karina", "Lucas", "Mariana", "Nicolas", "Olívia", "Pedro", "Quésia", "Rafael",
    "Sofia", "Thiago", "Úrsula", "Vinícius", "Yasmin", "Leonardo", "Beatriz", "Caio", "Débora",
    "Fernanda", "Gustavo", "Helena", "Igor", "Juliana", "Larissa", "Marcelo", "Natália",
    "Otávio", "Paula", "Renata", "Samuel", "Tatiane",
]
LAST_NAMES = [
    "Almeida", "Barros", "Cardoso", "Duarte", "Esteves", "Farias", "Gomes", "Henriques", "Lima",
    "Macedo", "Nogueira", "Oliveira", "Pacheco", "Queiroz", "Ribeiro", "Siqueira", "Teixeira",
    "Valente", "Xavier", "Zanetti", "Moura", "Prado", "Rezende", "Sampaio", "Tavares", "Vieira",
]

# Status VGCF no formato das planilhas. Sem interpretação aqui: a classificação
# (aprovado, reprovado, contratado...) é feita pelo administrador no sistema.
FINAL_STATUS_WEIGHTS = {
    "Pagamento Confirmado": 40, "Reprovado Entrevista": 14, "Cancelado VGCF": 13,
    "Aguardando Pagamento": 12, "Solicitado Kit Ferramenta": 6, "Enviado para MLPARCEIRO": 5,
    "Validação SUSEP": 4, "Entrevista METLIFE": 3, "Entrevista": 3,
}
RECENT_STATUS_WEIGHTS = {
    "Entrevista": 20, "Entrevista METLIFE": 15, "Validação SUSEP": 12, "Enviado para MLPARCEIRO": 12,
    "Aguardando Pagamento": 15, "Solicitado Kit Ferramenta": 6, "Pagamento Confirmado": 8,
    "Cancelado VGCF": 6, "Reprovado Entrevista": 6,
}
PAID = "Pagamento Confirmado"

# Eficiência diferente por agência, para que o ranking tenha o que mostrar.
_AGENCY_SKEW = {name: 0.6 + (i * 37 % 23) / 23 * 0.9 for i, (name, _) in enumerate(DEMO_AGENCIES)}


def _person(rng: random.Random) -> str:
    middle = rng.choice(LAST_NAMES)
    last = rng.choice([n for n in LAST_NAMES if n != middle])
    return f"{rng.choice(FIRST_NAMES)} {middle} {last}"


def _status(rng: random.Random, agency: str, recent: bool) -> str:
    weights = dict(RECENT_STATUS_WEIGHTS if recent else FINAL_STATUS_WEIGHTS)
    weights[PAID] = weights[PAID] * _AGENCY_SKEW[agency]
    codes = list(weights)
    return rng.choices(codes, weights=[weights[c] for c in codes])[0]


EXECUTIVES = ["AB", "CD", "EF", "GH"]
ENS_STATUS_WEIGHTS = {"aprovada": 60, "suspensa": 12, "recuperacao": 8, "reprovada": 8, "anulada": 6,
                      "nao avaliado": 6}


def fake_cpf(rng: random.Random) -> str:
    """CPF sorteado com dígitos verificadores válidos (fictício)."""
    digits = [rng.randint(0, 9) for _ in range(9)]
    for size in (9, 10):
        total = sum(d * (size + 1 - i) for i, d in enumerate(digits[:size]))
        digits.append((total * 10) % 11 % 10)
    cpf = "".join(map(str, digits))
    return f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"


def _branch(agency: str) -> str:
    return normalize_key(agency)[:3].upper()


def build_cohort_workbook(
    cohort: str,
    exam: date,
    *,
    size: int,
    seed: int,
    recent: bool = False,
) -> bytes:
    rng = random.Random(seed)
    sub_reps = {agency: [_person(rng) for _ in range(3)] for agency, _ in DEMO_AGENCIES}
    executives = {agency: EXECUTIVES[i % len(EXECUTIVES)] for i, (agency, _) in enumerate(DEMO_AGENCIES)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Base_selecao (demonstração)"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0061A0")

    evaluation = datetime(exam.year, exam.month, 1)
    for _ in range(size):
        agency, regional = rng.choice(DEMO_AGENCIES)
        status = _status(rng, agency, recent)
        suspended = rng.random() < 0.2
        ens = "suspensa" if suspended else rng.choices(
            list(ENS_STATUS_WEIGHTS), weights=list(ENS_STATUS_WEIGHTS.values())
        )[0]
        final = "cadastrada" if status == PAID or (not recent and rng.random() < 0.3) else "queda"
        registered = None
        if final == "cadastrada":
            month = exam.month + rng.randint(0, 3)
            registered = datetime(exam.year + (month - 1) // 12, (month - 1) % 12 + 1, 1)
        ws.append(
            [
                cohort,
                fake_cpf(rng),
                _person(rng),
                rng.choice(sub_reps[agency]),
                agency,
                f"{_branch(agency)} {rng.randint(1, 3)}" if rng.random() < 0.4 else _branch(agency),
                regional,
                executives[agency],
                "sim" if suspended else "nao",
                rng.choice([0, 1, 1, 1, 2]),
                status,
                ens,
                final,
                evaluation,
                registered,
            ]
        )
    for row in ws.iter_rows(min_row=2, min_col=14, max_col=15):
        for cell in row:
            cell.number_format = "mmm/yyyy"
    widths = (12, 16, 32, 28, 16, 10, 12, 10, 10, 8, 26, 14, 14, 14, 14)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def demo_cohorts(count: int = 8, today: date | None = None) -> list[tuple[str, bytes]]:
    """Turmas fictícias mensais, da mais antiga para a mais recente."""
    today = today or date.today()
    files = []
    for index in range(count):
        months_back = count - 1 - index
        exam = date(today.year, today.month, 20) - timedelta(days=30 * months_back)
        size = 45 + (index * 13) % 30
        name = f"DEMO T{50 + index}"
        content = build_cohort_workbook(name, exam, size=size, seed=index + 7, recent=months_back <= 1)
        files.append((f"{name.replace(' ', '_')}.xlsx", content))
    return files
