"""Leitura do modelo de planilha das turmas (ex.: "SAFEE T61 - Prova 21/01/26").

Estrutura observada no modelo recebido:
- linha 1: título mesclado com a turma e a data da prova;
- linha 2: cabeçalho com 10 colunas;
- a partir da linha 3: um candidato por linha, de várias agências.

O cabeçalho é localizado pelos nomes das colunas (sem diferenciar maiúsculas, acentos
ou pontuação), então pequenas variações de posição ou de grafia não quebram a leitura.
Colunas a mais são preservadas no registro original da linha.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.core.text import clean, normalize_key
from app.importing import vgcf

TEMPLATE_NAME = "turma-safee-v1"
LOCAL_TZ = ZoneInfo("America/Sao_Paulo")
HEADER_SEARCH_ROWS = 15


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    aliases: tuple[str, ...]
    required: bool = False


COLUMNS: tuple[Column, ...] = (
    Column("external_record_id", "ID", ("id", "id do registro"), required=True),
    Column(
        "site_registered_at",
        "Cadastro no site",
        ("cadastro no site", "data de cadastro", "data cadastro", "cadastro"),
    ),
    Column("agency", "Agência", ("agencia",), required=True),
    Column("regional", "Regional", ("regional",)),
    Column(
        "sub_representative",
        "Sub Representante",
        ("sub representante", "sub-representante", "subrepresentante"),
    ),
    Column(
        "full_name",
        "Nome Potencial Parceiro",
        ("nome potencial parceiro", "nome do candidato", "nome"),
        required=True,
    ),
    Column("susep_active", "SUSEP Ativa?", ("susep ativa", "susep ativa?")),
    Column("susep_number", "Número SUSEP", ("numero susep", "n susep", "no susep")),
    Column(
        "relationship_program",
        "Programa Relacionamento",
        ("programa relacionamento", "programa de relacionamento"),
    ),
    Column("status", "Status VGCF", ("status", "status vgcf"), required=True),
)

_HEADER_PUNCT = re.compile(r"[^a-z0-9 ]+")
_TITLE_EXAM = re.compile(r"prova\s*:?\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", re.IGNORECASE)
_TITLE_SPLIT = re.compile(r"\s+[-–—]\s+prova\b.*$", re.IGNORECASE)
_DATE_TEXT = re.compile(
    r"^(\d{1,2})/(\d{1,2})/(\d{2,4})(?:\s+(\d{1,2}):(\d{2})(?::(\d{2}))?)?$"
)


class TemplateError(Exception):
    """A planilha não segue o modelo esperado (lote inteiro rejeitado)."""


@dataclass
class Issue:
    code: str
    message: str
    field: str | None = None
    raw_value: str | None = None
    severity: str = "error"


@dataclass
class ParsedRow:
    row_number: int
    values: dict[str, object]
    raw: dict[str, str | None]
    issues: list[Issue] = field(default_factory=list)

    @property
    def rejected(self) -> bool:
        return any(i.severity == "error" for i in self.issues)


@dataclass
class ParsedSheet:
    sheet_name: str
    title: str | None
    cohort_name: str | None
    exam_date: date | None
    rows: list[ParsedRow]
    missing_optional: list[str]


def _header_key(value: object) -> str:
    return _HEADER_PUNCT.sub("", normalize_key(value)).strip()


def _match_columns(header: list[object]) -> dict[str, int]:
    found: dict[str, int] = {}
    keys = [_header_key(h) for h in header]
    for column in COLUMNS:
        aliases = {_header_key(a) for a in column.aliases}
        for index, key in enumerate(keys):
            if key and key in aliases and index not in found.values():
                found[column.key] = index
                break
    return found


def parse_title(title: str | None) -> tuple[str | None, date | None]:
    """ "SAFEE T61 - Prova 21/01/26" → ("SAFEE T61", 2026-01-21)."""
    if not title:
        return None, None
    exam: date | None = None
    match = _TITLE_EXAM.search(title)
    if match:
        day, month, year = (int(g) for g in match.groups())
        year = year + 2000 if year < 100 else year
        try:
            exam = date(year, month, day)
        except ValueError:
            exam = None
    cohort = clean(_TITLE_SPLIT.sub("", title))
    return cohort, exam


def parse_datetime(value: object) -> datetime | None:
    """Aceita data do Excel ou texto "dd/mm/aaaa hh:mm:ss" (horário de Brasília)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        local = value if value.tzinfo else value.replace(tzinfo=LOCAL_TZ)
        return local
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=LOCAL_TZ)
    match = _DATE_TEXT.match(str(value).strip())
    if not match:
        raise ValueError(str(value))
    day, month, year, hour, minute, second = match.groups()
    year_i = int(year) + 2000 if len(year) == 2 else int(year)
    return datetime(
        year_i,
        int(month),
        int(day),
        int(hour or 0),
        int(minute or 0),
        int(second or 0),
        tzinfo=LOCAL_TZ,
    )


def _as_identifier(value: object) -> str | None:
    """ID e números SUSEP: o Excel pode trazer 3094.0 ou texto; guardamos só o texto limpo."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def _cell_text(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def _find_header(ws) -> tuple[int, list[object], dict[str, int]] | None:
    for row_index, row in enumerate(
        ws.iter_rows(min_row=1, max_row=HEADER_SEARCH_ROWS, values_only=True), start=1
    ):
        header = list(row)
        found = _match_columns(header)
        if all(c.key in found for c in COLUMNS if c.required):
            return row_index, header, found
    return None


def parse_workbook(content: bytes) -> ParsedSheet:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # arquivo corrompido ou não-XLSX
        raise TemplateError("Não foi possível abrir o arquivo como planilha Excel (.xlsx).") from exc

    try:
        for ws in wb.worksheets:
            located = _find_header(ws)
            if located is None:
                continue
            header_row, header, columns = located

            title = None
            for row in ws.iter_rows(min_row=1, max_row=header_row - 1, values_only=True):
                title = next((clean(v) for v in row if clean(v)), None) or title
                if title:
                    break
            cohort, exam = parse_title(title)

            rows: list[ParsedRow] = []
            seen_ids: dict[str, int] = {}
            for row_number, row in enumerate(
                ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
            ):
                cells = list(row)
                if not any(clean(c) for c in cells):
                    continue
                parsed = _parse_row(row_number, cells, header, columns)
                record_id = parsed.values.get("external_record_id")
                if isinstance(record_id, str):
                    if record_id in seen_ids:
                        parsed.issues.append(
                            Issue(
                                "duplicate_id",
                                f"ID repetido na planilha (já aparece na linha {seen_ids[record_id]}).",
                                "external_record_id",
                                record_id,
                            )
                        )
                    else:
                        seen_ids[record_id] = row_number
                rows.append(parsed)

            missing = [c.label for c in COLUMNS if not c.required and c.key not in columns]
            return ParsedSheet(ws.title, title, cohort, exam, rows, missing)
    finally:
        wb.close()

    required = ", ".join(c.label for c in COLUMNS if c.required)
    raise TemplateError(
        f"Nenhuma aba tem o cabeçalho esperado. Colunas obrigatórias: {required}."
    )


def _parse_row(
    row_number: int, cells: list[object], header: list[object], columns: dict[str, int]
) -> ParsedRow:
    def cell(key: str) -> object:
        index = columns.get(key)
        return cells[index] if index is not None and index < len(cells) else None

    raw: dict[str, str | None] = {}
    for index, name in enumerate(header):
        label = clean(name) or f"Coluna {index + 1}"
        raw[label] = _cell_text(cells[index]) if index < len(cells) else None

    parsed = ParsedRow(row_number=row_number, values={}, raw=raw)
    values = parsed.values

    values["external_record_id"] = _as_identifier(cell("external_record_id"))
    values["agency"] = clean(cell("agency"))
    values["full_name"] = clean(cell("full_name"))
    for key, label in (
        ("external_record_id", "ID"),
        ("agency", "Agência"),
        ("full_name", "Nome"),
    ):
        if not values[key]:
            parsed.issues.append(Issue("required", f"{label} não preenchido.", key))

    values["regional"] = clean(cell("regional"))
    values["sub_representative"] = clean(cell("sub_representative"))
    values["susep_active"] = clean(cell("susep_active"))
    values["susep_number"] = _as_identifier(cell("susep_number"))
    values["relationship_program"] = clean(cell("relationship_program"))
    values["status"] = clean(cell("status"))
    if not values["status"]:
        parsed.issues.append(
            Issue("status_empty", "Status não preenchido.", "status", severity="warning")
        )
    elif vgcf.lookup(values["status"]) is None:
        parsed.issues.append(
            Issue(
                "status_unknown",
                "Status fora da lista oficial VGCF. Confira a grafia ou cadastre o novo status.",
                "status",
                values["status"],
                severity="warning",
            )
        )

    registered = cell("site_registered_at")
    try:
        values["site_registered_at"] = parse_datetime(registered)
    except ValueError:
        values["site_registered_at"] = None
        parsed.issues.append(
            Issue(
                "invalid_date",
                "Data de cadastro em formato não reconhecido (esperado dd/mm/aaaa hh:mm:ss).",
                "site_registered_at",
                _cell_text(registered),
                severity="warning",
            )
        )
    return parsed
