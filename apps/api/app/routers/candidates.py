"""CRM: lista de candidaturas, indicadores do topo, perfil do candidato, notas, exportação."""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Annotated, Literal

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import distinct, func, select
from sqlalchemy.orm import aliased

from app.auth.deps import DbSession, require_permission
from app.auth.permissions import Permission
from app.core.security import utcnow
from app.models import (
    UNMAPPED_OUTCOME,
    Agency,
    Application,
    ApplicationFieldChange,
    ApplicationStatusHistory,
    Candidate,
    ImportBatch,
    InternalNote,
    ProcessStatus,
    SavedFilter,
    User,
)
from app.schemas import Page
from app.services.audit import record_audit
from app.services.filters import OUTCOME_EXPR, ApplicationFilters, application_filters
from app.services.outcomes import outcome_codes, outcome_labels

router = APIRouter(prefix="/api", tags=["candidates"])

Viewer = Annotated[User, Depends(require_permission(Permission.CANDIDATES_VIEW))]
Exporter = Annotated[User, Depends(require_permission(Permission.CANDIDATES_EXPORT))]
NoteWriter = Annotated[User, Depends(require_permission(Permission.CANDIDATES_NOTES_WRITE))]
Filters = Annotated[ApplicationFilters, Depends(application_filters)]

SORTS = {
    "name": Candidate.full_name,
    "external_id": Application.external_record_id,
    "agency": Agency.name,
    "regional": Application.regional_norm,
    "registered_at": Application.site_registered_at,
    "status": Application.status_raw,
    "cohort": Application.cohort_name,
    "exam_date": Application.exam_date,
}


class StatusRef(BaseModel):
    id: int | None
    label: str | None
    outcome: str


class ApplicationRow(BaseModel):
    id: int
    external_record_id: str | None
    candidate_id: int
    full_name: str
    agency_id: int
    agency_name: str
    regional: str | None
    sub_representative: str | None
    cohort_name: str | None
    exam_date: date | None
    registered_at: datetime | None
    status_raw: str | None
    status: StatusRef
    susep_active: str | None
    susep_number: str | None
    relationship_program: str | None
    cpf_masked: str | None
    agency_branch: str | None
    executive: str | None
    suspended: str | None
    exams_count: int | None
    status_ens: str | None
    status_final: str | None
    evaluation_date: date | None
    is_demo: bool


class Summary(BaseModel):
    candidates: int
    applications: int
    by_outcome: dict[str, int]


def _base_query():
    return (
        select(Application, Candidate, Agency, ProcessStatus)
        .join(Candidate, Candidate.id == Application.candidate_id)
        .join(Agency, Agency.id == Application.agency_id)
        .outerjoin(ProcessStatus, ProcessStatus.id == Application.status_id)
    )


def _row(app: Application, cand: Candidate, agency: Agency, st: ProcessStatus | None) -> ApplicationRow:
    return ApplicationRow(
        id=app.id,
        external_record_id=app.external_record_id,
        candidate_id=cand.id,
        full_name=cand.full_name,
        agency_id=agency.id,
        agency_name=agency.name,
        regional=app.regional_norm,
        sub_representative=app.sub_representative_raw,
        cohort_name=app.cohort_name,
        exam_date=app.exam_date,
        registered_at=app.site_registered_at,
        status_raw=app.status_raw,
        status=StatusRef(
            id=st.id if st else None,
            label=st.label if st else app.status_raw,
            outcome=st.outcome if st else UNMAPPED_OUTCOME,
        ),
        susep_active=app.susep_active_raw,
        susep_number=app.susep_number,
        relationship_program=app.relationship_program_raw,
        cpf_masked=cand.cpf_masked,
        agency_branch=app.agency_branch_raw,
        executive=app.executive_raw,
        suspended=app.suspended_raw,
        exams_count=app.exams_count,
        status_ens=app.status_ens_raw,
        status_final=app.status_final_raw,
        evaluation_date=app.evaluation_date,
        is_demo=app.is_demo,
    )


@router.get("/candidates", response_model=Page[ApplicationRow])
def list_applications(
    _: Viewer,
    db: DbSession,
    filters: Filters,
    page: Annotated[int, Query(ge=1)] = 1,
    page_size: Annotated[int, Query(ge=1, le=200)] = 25,
    sort: Literal[
        "name", "external_id", "agency", "regional", "registered_at", "status", "cohort", "exam_date"
    ] = "registered_at",
    order: Literal["asc", "desc"] = "desc",
) -> Page[ApplicationRow]:
    query = filters.apply(_base_query())
    total = db.scalar(select(func.count()).select_from(query.with_only_columns(Application.id).subquery())) or 0
    column = SORTS[sort]
    ordering = column.desc().nulls_last() if order == "desc" else column.asc().nulls_last()
    rows = db.execute(
        query.order_by(ordering, Application.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return Page[ApplicationRow](
        items=[_row(*r) for r in rows], total=total, page=page, page_size=page_size
    )


@router.get("/candidates/summary", response_model=Summary)
def summary(_: Viewer, db: DbSession, filters: Filters) -> Summary:
    base = filters.apply(
        select(Application.id)
        .join(Candidate, Candidate.id == Application.candidate_id)
        .outerjoin(ProcessStatus, ProcessStatus.id == Application.status_id)
    )
    totals = db.execute(
        base.with_only_columns(
            func.count(Application.id), func.count(distinct(Application.candidate_id))
        )
    ).one()
    by_outcome = {o: 0 for o in outcome_codes(db)}
    for outcome, count in db.execute(
        base.with_only_columns(OUTCOME_EXPR, func.count(Application.id)).group_by(OUTCOME_EXPR)
    ):
        by_outcome[str(outcome)] = count
    return Summary(applications=totals[0], candidates=totals[1], by_outcome=by_outcome)


@router.get("/candidates/options")
def filter_options(_: Viewer, db: DbSession) -> dict:
    """Valores disponíveis para os filtros, com a contagem de candidaturas de cada um."""

    def distinct_values(column):
        rows = db.execute(
            select(column, func.count(Application.id))
            .where(column.is_not(None))
            .group_by(column)
            .order_by(column)
        )
        return [{"value": v, "count": c} for v, c in rows]

    agencies = db.execute(
        select(Agency.id, Agency.name, func.count(Application.id))
        .join(Application, Application.agency_id == Agency.id)
        .group_by(Agency.id, Agency.name)
        .order_by(Agency.name)
    )
    statuses = db.execute(
        select(ProcessStatus.id, ProcessStatus.label, ProcessStatus.outcome, func.count(Application.id))
        .join(Application, Application.status_id == ProcessStatus.id)
        .group_by(ProcessStatus.id, ProcessStatus.label, ProcessStatus.outcome)
        .order_by(ProcessStatus.label)
    )
    cohorts = db.execute(
        select(Application.cohort_name, func.max(Application.exam_date), func.count(Application.id))
        .where(Application.cohort_name.is_not(None))
        .group_by(Application.cohort_name)
        .order_by(func.max(Application.exam_date).desc())
    )
    sub_representatives = db.execute(
        select(Application.sub_representative_norm, Agency.name, func.count(Application.id))
        .join(Agency, Agency.id == Application.agency_id)
        .where(Application.sub_representative_norm.is_not(None))
        .group_by(Application.sub_representative_norm, Agency.name)
        .order_by(Application.sub_representative_norm, Agency.name)
    )
    reps: dict[str, dict] = {}
    for name, agency, count in sub_representatives:
        item = reps.setdefault(name, {"value": name, "agencies": [], "count": 0})
        item["agencies"].append(agency)
        item["count"] += count
    bounds = db.execute(select(func.min(Application.registered_at), func.max(Application.registered_at))).one()
    return {
        "agencies": [{"id": i, "name": n, "count": c} for i, n, c in agencies],
        "regionals": distinct_values(Application.regional_norm),
        "sub_representatives": list(reps.values()),
        "statuses": [
            {"id": i, "label": label, "outcome": o or UNMAPPED_OUTCOME, "count": c}
            for i, label, o, c in statuses
        ],
        "cohorts": [
            {"value": n, "exam_date": d.isoformat() if d else None, "count": c} for n, d, c in cohorts
        ],
        "programs": distinct_values(Application.relationship_program_raw),
        "susep_active": distinct_values(Application.susep_active_raw),
        "date_min": bounds[0].isoformat() if bounds[0] else None,
        "date_max": bounds[1].isoformat() if bounds[1] else None,
    }


# ---------------------------------------------------------------- perfil do candidato
class TimelineEvent(BaseModel):
    kind: Literal["created", "status", "field", "note"]
    at: datetime
    title: str
    detail: str | None = None
    author: str | None = None
    source: str | None = None


class ApplicationDetail(BaseModel):
    application: ApplicationRow
    raw_payload: dict | None
    first_import: dict | None
    other_applications: list[ApplicationRow]
    timeline: list[TimelineEvent]


@router.get("/applications/{application_id}", response_model=ApplicationDetail)
def application_detail(application_id: int, _: Viewer, db: DbSession) -> ApplicationDetail:
    found = db.execute(_base_query().where(Application.id == application_id)).first()
    if found is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Candidatura não encontrada.")
    app, cand, agency, st = found

    others = db.execute(
        _base_query()
        .where(Application.candidate_id == cand.id, Application.id != app.id)
        .order_by(Application.site_registered_at.desc())
    ).all()

    events: list[TimelineEvent] = []
    first_import = None
    if app.first_import_batch_id:
        batch = db.get(ImportBatch, app.first_import_batch_id)
        if batch:
            first_import = {"id": batch.id, "file_name": batch.file_name, "received_at": batch.received_at}

    from_status = aliased(ProcessStatus)
    to_status = aliased(ProcessStatus)
    for hist, from_label, to_label, batch_file in db.execute(
        select(ApplicationStatusHistory, from_status.label, to_status.label, ImportBatch.file_name)
        .outerjoin(from_status, from_status.id == ApplicationStatusHistory.from_status_id)
        .outerjoin(to_status, to_status.id == ApplicationStatusHistory.to_status_id)
        .outerjoin(ImportBatch, ImportBatch.id == ApplicationStatusHistory.import_batch_id)
        .where(ApplicationStatusHistory.application_id == app.id)
    ):
        new_label = to_label or hist.to_status_raw or "sem status"
        if hist.from_status_raw is None and hist.from_status_id is None:
            events.append(
                TimelineEvent(
                    kind="created",
                    at=hist.changed_at,
                    title=f"Entrada pela agência {agency.name} com status {new_label}",
                    detail=hist.note,
                    source=batch_file,
                )
            )
        else:
            events.append(
                TimelineEvent(
                    kind="status",
                    at=hist.changed_at,
                    title=f"Status alterado de {from_label or hist.from_status_raw} para {new_label}",
                    source=batch_file,
                )
            )
    for change, batch_file in db.execute(
        select(ApplicationFieldChange, ImportBatch.file_name)
        .outerjoin(ImportBatch, ImportBatch.id == ApplicationFieldChange.import_batch_id)
        .where(ApplicationFieldChange.application_id == app.id)
    ):
        events.append(
            TimelineEvent(
                kind="field",
                at=change.changed_at,
                title=f"{change.field} atualizado por importação",
                detail=f"{change.old_value or 'vazio'} → {change.new_value or 'vazio'}",
                source=batch_file,
            )
        )
    for note, author in db.execute(
        select(InternalNote, User.name)
        .outerjoin(User, User.id == InternalNote.author_id)
        .where(InternalNote.application_id == app.id)
    ):
        events.append(
            TimelineEvent(kind="note", at=note.created_at, title="Observação interna", detail=note.body, author=author)
        )
    events.sort(key=lambda e: (e.at.replace(tzinfo=None) if e.at.tzinfo else e.at), reverse=True)

    return ApplicationDetail(
        application=_row(app, cand, agency, st),
        raw_payload=app.raw_payload,
        first_import=first_import,
        other_applications=[_row(*r) for r in others],
        timeline=events,
    )


class NoteIn(BaseModel):
    body: str = Field(min_length=1, max_length=4000)


@router.post("/applications/{application_id}/notes", status_code=status.HTTP_201_CREATED)
def add_note(application_id: int, payload: NoteIn, request: Request, user: NoteWriter, db: DbSession) -> dict:
    if db.get(Application, application_id) is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Candidatura não encontrada.")
    note = InternalNote(application_id=application_id, author_id=user.id, body=payload.body.strip())
    db.add(note)
    db.flush()
    record_audit(
        db,
        action="candidates.note_added",
        request=request,
        actor=user,
        entity_type="application",
        entity_id=application_id,
    )
    return {"id": note.id}


# ---------------------------------------------------------------- exportação
EXPORT_LIMIT = 50_000


@router.get("/candidates/export")
def export_applications(request: Request, user: Exporter, db: DbSession, filters: Filters):
    rows = db.execute(
        filters.apply(_base_query()).order_by(Application.site_registered_at.desc()).limit(EXPORT_LIMIT)
    ).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidaturas"
    header = [
        "Turma", "CPF (mascarado)", "Candidato", "R.N", "Nome Agência", "Agência", "Regional", "Executivo",
        "Suspenso", "Provas", "Status VGCF", "Classificação", "Status ENS", "Status Final", "Data Avaliação",
        "Data Cadastro", "Data da prova", "ID", "SUSEP Ativa?", "Número SUSEP", "Programa Relacionamento",
    ]
    ws.append(header)
    labels = outcome_labels(db)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0061A0")
    for r in rows:
        row = _row(*r)
        ws.append(
            [
                row.cohort_name,
                row.cpf_masked,
                row.full_name,
                row.sub_representative,
                row.agency_name,
                row.agency_branch,
                row.regional,
                row.executive,
                row.suspended,
                row.exams_count,
                row.status_raw,
                labels.get(row.status.outcome, row.status.outcome),
                row.status_ens,
                row.status_final,
                row.evaluation_date,
                row.registered_at.replace(tzinfo=None) if row.registered_at else None,
                row.exam_date,
                row.external_record_id,
                row.susep_active,
                row.susep_number,
                row.relationship_program,
            ]
        )
    ws.freeze_panes = "A2"
    stamp = utcnow()
    meta = wb.create_sheet("Exportação")
    meta.append(["Exportado por", user.email])
    meta.append(["Data (UTC)", stamp.strftime("%d/%m/%Y %H:%M")])
    meta.append(["Registros", len(rows)])
    meta.append(["Filtros", str(filters.as_dict() or "nenhum")])
    meta.append(["Aviso", "Contém dados pessoais. Uso interno; não compartilhe fora da equipe."])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    record_audit(
        db,
        action="candidates.exported",
        request=request,
        actor=user,
        details={"rows": len(rows), "filters": filters.as_dict()},
    )
    filename = f"candidaturas_{stamp.strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------- filtros salvos
class SavedFilterIn(BaseModel):
    scope: Literal["candidates", "performance"]
    name: str = Field(min_length=1, max_length=120)
    definition: dict
    is_shared: bool = False


class SavedFilterOut(BaseModel):
    id: int
    scope: str
    name: str
    definition: dict
    is_shared: bool
    is_mine: bool


@router.get("/saved-filters", response_model=list[SavedFilterOut])
def list_saved_filters(
    user: Viewer, db: DbSession, scope: Literal["candidates", "performance"] = "candidates"
) -> list[SavedFilterOut]:
    rows = db.scalars(
        select(SavedFilter)
        .where(SavedFilter.scope == scope)
        .where((SavedFilter.user_id == user.id) | SavedFilter.is_shared)
        .order_by(SavedFilter.name)
    )
    return [
        SavedFilterOut(
            id=f.id, scope=f.scope, name=f.name, definition=f.definition,
            is_shared=f.is_shared, is_mine=f.user_id == user.id,
        )
        for f in rows
    ]


@router.post("/saved-filters", response_model=SavedFilterOut, status_code=status.HTTP_201_CREATED)
def save_filter(payload: SavedFilterIn, user: Viewer, db: DbSession) -> SavedFilterOut:
    existing = db.scalar(
        select(SavedFilter).where(
            SavedFilter.user_id == user.id,
            SavedFilter.scope == payload.scope,
            SavedFilter.name == payload.name.strip(),
        )
    )
    saved = existing or SavedFilter(user_id=user.id, scope=payload.scope, name=payload.name.strip())
    saved.definition = payload.definition
    saved.is_shared = payload.is_shared
    db.add(saved)
    db.commit()
    return SavedFilterOut(
        id=saved.id, scope=saved.scope, name=saved.name, definition=saved.definition,
        is_shared=saved.is_shared, is_mine=True,
    )


@router.delete("/saved-filters/{filter_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_saved_filter(filter_id: int, user: Viewer, db: DbSession) -> None:
    saved = db.get(SavedFilter, filter_id)
    if saved is None or saved.user_id != user.id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Filtro não encontrado.")
    db.delete(saved)
    db.commit()
