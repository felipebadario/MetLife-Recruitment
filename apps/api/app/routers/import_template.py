"""Padrão da planilha: quais colunas a importação lê, com que nomes e quais são obrigatórias."""

from __future__ import annotations

import io
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Request, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import delete, select

from app.auth.deps import DbSession, require_permission
from app.auth.permissions import Permission
from app.importing.columns import NAME_MAX, check_columns, load_columns
from app.importing.template import FIELDS, MAX_CUSTOM_FIELDS, Column, is_custom
from app.models import ImportColumn, User
from app.services.audit import record_audit

router = APIRouter(prefix="/api/admin/import-template", tags=["catalogs"])

CatalogManager = Annotated[User, Depends(require_permission(Permission.CATALOGS_MANAGE))]


class ColumnOut(BaseModel):
    field: str
    name: str
    description: str
    kind: str
    headers: list[str]
    required: bool
    active: bool
    locked: bool
    custom: bool
    default_name: str | None
    default_headers: list[str]
    default_required: bool
    default_active: bool


class TemplateOut(BaseModel):
    columns: list[ColumnOut]
    is_default: bool


class ColumnIn(BaseModel):
    field: str | None = Field(default=None, max_length=40)  # vazio: campo novo
    name: str | None = Field(default=None, max_length=NAME_MAX)
    headers: list[Annotated[str, Field(max_length=120)]] = Field(max_length=10)
    required: bool = False
    active: bool = True


class TemplateIn(BaseModel):
    columns: list[ColumnIn] = Field(max_length=len(FIELDS) + MAX_CUSTOM_FIELDS)


def _is_default(column: Column) -> bool:
    spec = column.spec
    return spec is not None and (column.display_name, list(column.headers), column.required, column.active) == (
        spec.name, list(spec.headers), spec.required, spec.active
    )


def _template(db) -> TemplateOut:
    columns = load_columns(db)
    out = []
    for column in columns:
        spec = column.spec
        out.append(
            ColumnOut(
                field=column.key, name=column.display_name, kind=column.kind,
                description=spec.description if spec else "Campo criado aqui. Guardado como veio na planilha.",
                headers=list(column.headers), required=column.required, active=column.active,
                locked=bool(spec and spec.locked), custom=spec is None,
                default_name=spec.name if spec else None,
                default_headers=list(spec.headers) if spec else [],
                default_required=spec.required if spec else False,
                default_active=spec.active if spec else True,
            )
        )
    return TemplateOut(columns=out, is_default=all(_is_default(c) for c in columns))


def _snapshot(columns: list[Column]) -> dict[str, dict]:
    return {
        c.key: {"nome": c.display_name, "headers": list(c.headers), "required": c.required, "active": c.active}
        for c in columns
    }


@router.get("", response_model=TemplateOut)
def get_template(_: CatalogManager, db: DbSession) -> TemplateOut:
    return _template(db)


@router.put("", response_model=TemplateOut)
def save_template(payload: TemplateIn, request: Request, user: CatalogManager, db: DbSession) -> TemplateOut:
    before = _snapshot(load_columns(db))
    saved = {row.field: row for row in db.scalars(select(ImportColumn))}
    try:
        columns = check_columns(
            [Column(c.field or "", tuple(c.headers), c.required, c.active, c.name or "") for c in payload.columns],
            existing={key for key in saved if is_custom(key)},
        )
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)) from exc

    keys = {c.key for c in columns}
    for key, row in saved.items():
        if is_custom(key) and key not in keys:
            db.delete(row)
    for column in columns:
        row = saved.get(column.key) or ImportColumn(field=column.key)
        spec = column.spec
        row.label = None if spec is not None and column.name == spec.name else column.name
        row.headers = list(column.headers)
        row.is_required = column.required
        row.is_active = column.active
        db.add(row)
    db.flush()

    after = _snapshot(load_columns(db))
    changes = {
        (value or before[key])["nome"]: {"antes": before.get(key), "depois": value}
        for key in before.keys() | after.keys()
        if before.get(key) != (value := after.get(key))
    }
    if changes:
        record_audit(
            db, action="catalog.import_template_updated", request=request, actor=user,
            entity_type="import_template", details={"changes": changes},
        )
    else:
        db.commit()
    return _template(db)


@router.delete("", response_model=TemplateOut)
def reset_template(request: Request, user: CatalogManager, db: DbSession) -> TemplateOut:
    """Volta ao padrão inicial (planilha "Base_selecao"). Os campos criados são retirados."""
    db.execute(delete(ImportColumn))
    db.flush()
    record_audit(
        db, action="catalog.import_template_reset", request=request, actor=user, entity_type="import_template",
    )
    return _template(db)


@router.get("/sample")
def template_sample(_: CatalogManager, db: DbSession) -> Response:
    """Planilha vazia só com o cabeçalho do padrão em vigor."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Modelo"
    ws.append([c.label for c in load_columns(db) if c.active])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0061A0")
        ws.column_dimensions[cell.column_letter].width = max(12, len(str(cell.value)) + 4)
    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo_planilha_turma.xlsx"'},
    )
