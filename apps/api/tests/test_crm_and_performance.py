"""CRM, catálogos e dashboard sobre dados fictícios importados."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import load_workbook

from app.importing.demo import build_cohort_workbook
from app.models import UserRole
from tests.conftest import login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _seed(client, make_user, size: int = 40):
    make_user("admin@example.com", UserRole.ADMIN)
    login(client, "admin@example.com")
    content = build_cohort_workbook("DEMO T80", date(2026, 3, 10), first_id=7000, size=size, seed=3)
    client.post("/api/imports", files={"file": ("t80.xlsx", content, XLSX)})


def _classify(client, mapping: dict[str, str]):
    for status in client.get("/api/admin/statuses").json():
        if status["label"] in mapping:
            response = client.patch(
                f"/api/admin/statuses/{status['id']}", json={"outcome": mapping[status["label"]]}
            )
            assert response.status_code == 200


def test_candidates_list_filters_sort_and_summary(client, make_user):
    _seed(client, make_user)
    page = client.get("/api/candidates", params={"page_size": 10, "sort": "name", "order": "asc"}).json()
    assert page["total"] == 40 and len(page["items"]) == 10
    names = [r["full_name"] for r in page["items"]]
    assert names == sorted(names)

    options = client.get("/api/candidates/options").json()
    agency = options["agencies"][0]
    filtered = client.get("/api/candidates", params={"agency_id": agency["id"]}).json()
    assert filtered["total"] == agency["count"]
    assert all(r["agency_id"] == agency["id"] for r in filtered["items"])

    summary = client.get("/api/candidates/summary").json()
    assert summary["applications"] == 40
    assert summary["by_outcome"]["unmapped"] == 40

    first = page["items"][0]
    search = client.get("/api/candidates", params={"search": first["full_name"][:6].lower()}).json()
    assert any(r["id"] == first["id"] for r in search["items"])


def test_classifying_statuses_drives_rates(client, make_user):
    _seed(client, make_user)
    _classify(client, {
        "Pagamento Confirmado": "hired", "Reprovado Entrevista": "rejected", "Cancelado VGCF": "withdrawn",
        "Aguardando Pagamento": "in_progress", "Enviado para MLPARCEIRO": "in_progress",
        "Entrevista": "in_progress", "Entrevista METLIFE": "in_progress", "Validação SUSEP": "in_progress",
        "Solicitado Kit Ferramenta": "in_progress",
    })
    data = client.get("/api/performance/dashboard").json()
    assert data["totals"]["applications"] == 40
    assert data["totals"]["by_outcome"]["unmapped"] == 0
    hiring = next(m for m in data["metrics"] if m["code"] == "hiring_rate")
    assert hiring["denominator"] == 40
    assert hiring["numerator"] == data["totals"]["by_outcome"]["hired"]
    assert sum(a["total"] for a in data["agencies"]) == 40
    assert data["timeline"] and data["cohorts"][0]["name"] == "DEMO T80"

    summary = client.get("/api/candidates/summary", params={"outcome": "hired"}).json()
    assert summary["applications"] == hiring["numerator"]


def test_dashboard_lists_unmapped_statuses(client, make_user):
    _seed(client, make_user)
    data = client.get("/api/performance/dashboard").json()
    assert data["unmapped_statuses"]
    assert all(m["value"] is None for m in data["metrics"])


def test_funnel_uses_stage_positions(client, make_user):
    _seed(client, make_user)
    stages = client.get("/api/admin/funnel-stages").json()
    last = stages[-1]
    for status in client.get("/api/admin/statuses").json():
        if status["label"] == "Pagamento Confirmado":
            client.patch(f"/api/admin/statuses/{status['id']}", json={"funnel_stage_id": last["id"]})
    data = client.get("/api/performance/dashboard").json()
    funnel = {f["label"]: f["count"] for f in data["funnel"]}
    paga = next(s["count"] for s in data["statuses"] if s["label"] == "Pagamento Confirmado")
    assert funnel[last["label"]] == paga
    assert funnel[stages[0]["label"]] == paga  # quem chegou ao fim passou pelas etapas anteriores


def test_metric_formula_is_configurable(client, make_user):
    _seed(client, make_user)
    response = client.patch(
        "/api/admin/metrics/approval_rate",
        json={"numerator_outcomes": ["approved"], "denominator_outcomes": ["approved", "rejected"]},
    )
    assert response.status_code == 200
    assert response.json()["denominator_outcomes"] == ["approved", "rejected"]
    bad = client.patch("/api/admin/metrics/approval_rate", json={"denominator_outcomes": []})
    assert bad.status_code == 422


def test_detail_notes_and_export(client, make_user):
    _seed(client, make_user)
    app_id = client.get("/api/candidates").json()["items"][0]["id"]
    assert client.post(f"/api/applications/{app_id}/notes", json={"body": "Retornar contato"}).status_code == 201
    detail = client.get(f"/api/applications/{app_id}").json()
    kinds = {e["kind"] for e in detail["timeline"]}
    assert {"created", "note"} <= kinds
    assert detail["raw_payload"]["Agência"] == detail["application"]["agency_name"]

    export = client.get("/api/candidates/export")
    assert export.status_code == 200
    wb = load_workbook(io.BytesIO(export.content))
    assert wb["Candidaturas"].max_row == 41


def test_saved_filters(client, make_user):
    _seed(client, make_user)
    created = client.post(
        "/api/saved-filters",
        json={"scope": "candidates", "name": "Só SUD", "definition": {"regional": ["SUD"]}},
    ).json()
    assert [f["name"] for f in client.get("/api/saved-filters").json()] == ["Só SUD"]
    assert client.delete(f"/api/saved-filters/{created['id']}").status_code == 204


def test_viewer_permissions(client, make_user):
    _seed(client, make_user)
    client.post("/api/auth/logout")
    make_user("viewer@example.com")
    login(client, "viewer@example.com")
    assert client.get("/api/candidates").status_code == 200
    app_id = client.get("/api/candidates").json()["items"][0]["id"]
    assert client.get(f"/api/applications/{app_id}").status_code == 200
    assert client.get("/api/candidates/export").status_code == 403
    assert client.post(f"/api/applications/{app_id}/notes", json={"body": "x"}).status_code == 403
    assert client.get("/api/performance/dashboard").status_code == 403
    assert client.get("/api/admin/statuses").status_code == 403
