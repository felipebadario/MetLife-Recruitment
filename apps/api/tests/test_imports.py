"""Importação no modelo real da planilha (dados 100% fictícios)."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db.session import get_sessionmaker
from app.importing.demo import HEADER, build_cohort_workbook
from app.importing.template import parse_title, parse_workbook
from app.models import (
    Application,
    ApplicationStatusHistory,
    Candidate,
    ProcessOutcome,
    ProcessStatus,
    UserRole,
)
from tests.conftest import login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook(size: int = 20, seed: int = 1) -> bytes:
    return build_cohort_workbook("DEMO T70", date(2026, 1, 21), first_id=5000, size=size, seed=seed)


def _modify(content: bytes, row: int, column: int, value: object) -> bytes:
    wb = load_workbook(io.BytesIO(content))
    wb.active.cell(row=row, column=column).value = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _upload(client, content: bytes, name: str = "turma.xlsx"):
    return client.post("/api/imports", files={"file": (name, content, XLSX)})


def _admin(client, make_user):
    make_user("admin@example.com", UserRole.ADMIN)
    login(client, "admin@example.com")


def test_title_parsing():
    assert parse_title("SAFEE T61 - Prova 21/01/26") == ("SAFEE T61", date(2026, 1, 21))
    assert parse_title("Turma sem data") == ("Turma sem data", None)


def test_parser_reads_template():
    sheet = parse_workbook(_workbook(size=12))
    assert sheet.cohort_name == "DEMO T70"
    assert sheet.exam_date == date(2026, 1, 21)
    assert len(sheet.rows) == 12
    assert not any(r.rejected for r in sheet.rows)
    assert list(sheet.rows[0].raw) == HEADER


def test_import_creates_records_and_unmapped_statuses(client, make_user):
    _admin(client, make_user)
    response = _upload(client, _workbook())
    assert response.status_code == 202

    batch = client.get(f"/api/imports/{response.json()['id']}").json()["batch"]
    assert batch["status"] == "completed"
    assert batch["rows_created"] == 20 and batch["rows_rejected"] == 0
    assert batch["cohort_name"] == "DEMO T70"

    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).select_from(Application)) == 20
        statuses = list(db.scalars(select(ProcessStatus)))
        # Nenhum status é interpretado automaticamente.
        assert statuses and all(s.outcome == ProcessOutcome.UNMAPPED for s in statuses)


def test_same_file_twice_is_skipped(client, make_user):
    _admin(client, make_user)
    content = _workbook()
    _upload(client, content)
    second = _upload(client, content).json()
    assert second["status"] == "skipped_duplicate"
    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).select_from(Application)) == 20


def test_reprocess_is_idempotent(client, make_user):
    _admin(client, make_user)
    first = _upload(client, _workbook()).json()
    again = client.post(f"/api/imports/{first['id']}/reprocess").json()
    detail = client.get(f"/api/imports/{again['id']}").json()["batch"]
    assert detail["rows_unchanged"] == 20 and detail["rows_created"] == 0


def test_status_change_is_recorded_in_history(client, make_user):
    _admin(client, make_user)
    content = _workbook()
    _upload(client, content)
    # Linha 3 = primeiro candidato; coluna 10 = Status
    updated = _modify(content, row=3, column=10, value="NOVO-STATUS")
    batch = _upload(client, updated, "turma_atualizada.xlsx").json()
    detail = client.get(f"/api/imports/{batch['id']}").json()["batch"]
    assert detail["rows_updated"] == 1 and detail["rows_unchanged"] == 19

    with get_sessionmaker()() as db:
        app = db.scalar(select(Application).where(Application.external_record_id == "5000"))
        assert app.status_raw == "NOVO-STATUS"
        history = list(
            db.scalars(
                select(ApplicationStatusHistory).where(ApplicationStatusHistory.application_id == app.id)
            )
        )
        assert len(history) == 2

    timeline = client.get(f"/api/applications/{app.id}").json()["timeline"]
    assert any(e["kind"] == "status" for e in timeline)


def test_invalid_rows_are_rejected_and_reported(client, make_user):
    _admin(client, make_user)
    content = _modify(_workbook(size=5), row=4, column=6, value=None)  # sem nome
    content = _modify(content, row=5, column=1, value=5000)  # ID repetido
    batch = _upload(client, content).json()
    detail = client.get(f"/api/imports/{batch['id']}").json()
    assert detail["batch"]["status"] == "completed_with_errors"
    assert detail["batch"]["rows_rejected"] == 2
    codes = {i["code"] for i in detail["issues"]}
    assert {"required", "duplicate_id"} <= codes


def test_wrong_template_fails_whole_batch(client, make_user):
    _admin(client, make_user)
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["Coluna A", "Coluna B"])
    out = io.BytesIO()
    wb.save(out)
    batch = _upload(client, out.getvalue()).json()
    detail = client.get(f"/api/imports/{batch['id']}").json()["batch"]
    assert detail["status"] == "failed"
    assert "cabeçalho" in detail["error_message"]


def test_non_xlsx_rejected(client, make_user):
    _admin(client, make_user)
    response = client.post("/api/imports", files={"file": ("dados.csv", b"a,b", "text/csv")})
    assert response.status_code == 422


def test_same_person_in_two_cohorts_is_one_candidate(client, make_user):
    _admin(client, make_user)
    first = _workbook(size=3, seed=5)
    _upload(client, first)
    sheet = parse_workbook(first)
    name = sheet.rows[0].values["full_name"]
    second = build_cohort_workbook("DEMO T71", date(2026, 2, 21), first_id=6000, size=3, seed=6)
    second = _modify(second, row=3, column=6, value=name)
    _upload(client, second, "t71.xlsx")
    with get_sessionmaker()() as db:
        candidate = db.scalar(select(Candidate).where(Candidate.full_name == name))
        apps = db.scalar(select(func.count()).where(Application.candidate_id == candidate.id))
        assert apps == 2


def test_viewer_cannot_import(client, make_user):
    make_user("v@example.com")
    login(client, "v@example.com")
    assert _upload(client, _workbook()).status_code == 403


def test_operational_can_import_but_not_reset(client, make_user):
    make_user("op@example.com", UserRole.OPERATIONAL)
    login(client, "op@example.com")
    batch = _upload(client, _workbook(size=5)).json()
    assert client.get(f"/api/imports/{batch['id']}").json()["batch"]["status"] == "completed"
    assert client.get("/api/imports/sample").status_code == 200
    assert client.post(f"/api/imports/{batch['id']}/reprocess").status_code == 202
    assert client.delete(f"/api/imports/{batch['id']}").status_code == 200
    assert client.post("/api/imports/reset", json={"confirmation": "APAGAR"}).status_code == 403


def test_reset_requires_confirmation(client, make_user):
    _admin(client, make_user)
    _upload(client, _workbook())
    assert client.post("/api/imports/reset", json={"confirmation": "sim"}).status_code == 422
    assert client.post("/api/imports/reset", json={"confirmation": "APAGAR"}).status_code == 204
    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).select_from(Application)) == 0


def test_sample_download(client, make_user):
    _admin(client, make_user)
    response = client.get("/api/imports/sample")
    assert response.status_code == 200
    assert parse_workbook(response.content).cohort_name == "DEMO T99"


def test_vgcf_statuses_are_listed_before_any_import(client, make_user):
    _admin(client, make_user)
    labels = [s["label"] for s in client.get("/api/admin/statuses").json()]
    assert labels == [
        "Aguardando Pagamento", "Enviado para MLPARCEIRO", "Cancelado VGCF", "Reprovado Entrevista",
        "Pagamento Confirmado", "Entrevista METLIFE", "Solicitado Kit Ferramenta", "Entrevista",
        "Validação SUSEP",
    ]
    assert all(s["outcome"] == "unmapped" for s in client.get("/api/admin/statuses").json())


def test_vgcf_full_names_share_the_official_status(client, make_user):
    _admin(client, make_user)
    content = _modify(_workbook(size=4), row=3, column=10, value="Pagamento Confirmado")
    content = _modify(content, row=4, column=10, value="  pagamento   confirmado ")
    content = _modify(content, row=5, column=10, value="Validacao SUSEP")
    content = _modify(content, row=6, column=10, value="PAGA")  # abreviação não é reconhecida
    batch = _upload(client, content).json()
    detail = client.get(f"/api/imports/{batch['id']}").json()
    unknown = [i for i in detail["issues"] if i["code"] == "status_unknown"]
    assert [i["raw_value"] for i in unknown] == ["PAGA"]

    statuses = {s["label"]: s for s in client.get("/api/admin/statuses").json()}
    assert statuses["Pagamento Confirmado"]["applications"] == 2
    assert statuses["Validação SUSEP"]["applications"] == 1
    assert statuses["PAGA"]["applications"] == 1
    assert not statuses["PAGA"]["is_official"] and statuses["Entrevista"]["is_official"]
    assert list(statuses)[-1] == "PAGA"  # fora da lista oficial vai para o fim


def test_status_abbreviation_is_display_only(client, make_user):
    _admin(client, make_user)
    status = next(s for s in client.get("/api/admin/statuses").json() if s["label"] == "Pagamento Confirmado")
    response = client.patch(f"/api/admin/statuses/{status['id']}", json={"abbreviation": " PAGA "})
    assert response.status_code == 200 and response.json()["abbreviation"] == "PAGA"
    assert response.json()["outcome"] == "unmapped"

    cleared = client.patch(f"/api/admin/statuses/{status['id']}", json={"abbreviation": ""}).json()
    assert cleared["abbreviation"] is None


def _counts():
    with get_sessionmaker()() as db:
        return (
            db.scalar(select(func.count()).select_from(Application)),
            db.scalar(select(func.count()).select_from(Candidate)),
        )


def test_delete_batch_removes_only_its_applications(client, make_user):
    _admin(client, make_user)
    first = _upload(client, _workbook(size=10)).json()
    other = build_cohort_workbook("DEMO T71", date(2026, 2, 10), first_id=6000, size=5, seed=7)
    _upload(client, other, "t71.xlsx")
    assert _counts()[0] == 15

    response = client.delete(f"/api/imports/{first['id']}")
    assert response.status_code == 200
    assert response.json()["applications_removed"] == 10
    assert _counts()[0] == 5
    assert client.get(f"/api/imports/{first['id']}").status_code == 404


def test_delete_duplicate_restores_previous_data(client, make_user):
    _admin(client, make_user)
    content = _workbook(size=5)
    _upload(client, content)
    original = _status_of("5000")
    wrong_status = "Entrevista" if original != "Entrevista" else "Validação SUSEP"
    # A mesma turma enviada de novo, com um status diferente para o primeiro candidato.
    duplicate = _upload(client, _modify(content, row=3, column=10, value=wrong_status), "duplicada.xlsx").json()
    assert _status_of("5000") == wrong_status

    result = client.delete(f"/api/imports/{duplicate['id']}").json()
    assert result == {"applications_removed": 0, "applications_restored": 1, "candidates_removed": 0}
    assert _status_of("5000") == original
    assert _counts()[0] == 5
    with get_sessionmaker()() as db:
        app = db.scalar(select(Application).where(Application.external_record_id == "5000"))
        history = db.scalar(
            select(func.count()).where(ApplicationStatusHistory.application_id == app.id)
        )
        assert history == 1  # a mudança trazida pela planilha excluída some do histórico


def _status_of(record_id: str) -> str:
    with get_sessionmaker()() as db:
        return db.scalar(select(Application.status_raw).where(Application.external_record_id == record_id))


def test_viewer_cannot_delete_batch(client, make_user):
    _admin(client, make_user)
    batch = _upload(client, _workbook(size=3)).json()
    make_user("viewer@example.com", UserRole.VIEWER)
    login(client, "viewer@example.com")
    assert client.delete(f"/api/imports/{batch['id']}").status_code == 403


def test_bulk_upload_processes_in_order(client, make_user):
    _admin(client, make_user)
    files = [
        ("files", (f"t{n}.xlsx", build_cohort_workbook(f"DEMO T{n}", date(2026, 1, n), first_id=n * 100,
                                                      size=3, seed=n), XLSX))
        for n in range(1, 4)
    ]
    files.append(("files", ("t1-de-novo.xlsx", files[0][1][1], XLSX)))
    files.append(("files", ("notas.txt", b"x", "text/plain")))
    response = client.post("/api/imports/bulk", files=files)
    assert response.status_code == 202
    results = response.json()
    assert [bool(r["batch"]) for r in results] == [True, True, True, False, False]
    assert "repetido" in results[3]["error"] and ".xlsx" in results[4]["error"]
    batches = client.get("/api/imports").json()["items"]
    assert len(batches) == 3 and all(b["status"] == "completed" for b in batches)
    assert _counts()[0] == 9


def test_bulk_upload_limit(client, make_user):
    _admin(client, make_user)
    content = _workbook(size=1)
    files = [("files", (f"t{n}.xlsx", content, XLSX)) for n in range(21)]
    response = client.post("/api/imports/bulk", files=files)
    assert response.status_code == 422
    assert "20" in response.json()["detail"]
