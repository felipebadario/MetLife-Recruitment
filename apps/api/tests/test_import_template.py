"""Padrão da planilha configurável em Administração (dados 100% fictícios)."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db.session import get_sessionmaker
from app.importing.demo import HEADER, build_cohort_workbook
from app.models import Application, AuditLog, UserRole
from tests.conftest import login

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
URL = "/api/admin/import-template"


def _admin(client, make_user):
    make_user("admin@example.com", UserRole.ADMIN)
    login(client, "admin@example.com")


def _rename_header(content: bytes, old: str, new: str | None, *, title: str | None = None) -> bytes:
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    column = HEADER.index(old) + 1
    ws.cell(row=1, column=column).value = new
    if title:
        ws.insert_rows(1)
        ws.cell(row=1, column=1).value = title
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _with(columns: list[dict], field: str, **changes) -> list[dict]:
    return [{**c, **changes} if c["field"] == field else c for c in columns]


def _payload(columns: list[dict]) -> dict:
    return {
        "columns": [
            {k: c[k] for k in ("field", "headers", "required", "active")} for c in columns
        ]
    }


def _import(client, content: bytes) -> dict:
    batch = client.post("/api/imports", files={"file": ("turma.xlsx", content, XLSX)}).json()
    return client.get(f"/api/imports/{batch['id']}").json()


def test_default_template_is_base_selecao(client, make_user):
    _admin(client, make_user)
    data = client.get(URL).json()
    assert data["is_default"]
    active = [c["headers"][0] for c in data["columns"] if c["active"]]
    assert active == HEADER
    required = {c["headers"][0] for c in data["columns"] if c["required"]}
    assert required == {"Candidato", "Nome Agência", "Status VGCF"}


def test_renamed_column_is_read_after_saving(client, make_user):
    _admin(client, make_user)
    content = build_cohort_workbook("DEMO T90", date(2026, 4, 1), size=4, seed=2)
    renamed = _rename_header(content, "Candidato", "Nome completo")
    assert _import(client, renamed)["batch"]["status"] == "failed"

    columns = client.get(URL).json()["columns"]
    response = client.put(URL, json=_payload(_with(columns, "full_name", headers=["Nome completo", "Candidato"])))
    assert response.status_code == 200 and not response.json()["is_default"]

    detail = _import(client, renamed)
    assert detail["batch"]["status"] == "completed" and detail["batch"]["rows_created"] == 4
    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).where(AuditLog.action == "catalog.import_template_updated")) == 1


def test_inactive_column_is_ignored_and_title_gives_cohort(client, make_user):
    _admin(client, make_user)
    columns = client.get(URL).json()["columns"]
    columns = _with(columns, "cohort_name", active=False)
    columns = _with(columns, "executive", required=True)
    assert client.put(URL, json=_payload(columns)).status_code == 200

    content = build_cohort_workbook("DEMO T91", date(2026, 4, 1), size=3, seed=4)
    content = _rename_header(content, "Turma", "Outra coisa", title="SAFEE T91 - Prova 21/04/26")
    detail = _import(client, content)
    assert detail["batch"]["cohort_name"] == "SAFEE T91"
    assert detail["batch"]["exam_date"] == "2026-04-21"
    with get_sessionmaker()() as db:
        assert set(db.scalars(select(Application.cohort_name))) == {"SAFEE T91"}

    # Coluna marcada como obrigatória que não está na planilha: o arquivo não é importado.
    missing = _rename_header(build_cohort_workbook("DEMO T92", date(2026, 4, 1), size=2, seed=5), "Executivo", None)
    detail = _import(client, missing)
    assert detail["batch"]["status"] == "failed" and "Executivo" in detail["batch"]["error_message"]


def test_invalid_templates_are_rejected(client, make_user):
    _admin(client, make_user)
    columns = client.get(URL).json()["columns"]
    duplicated = _with(columns, "executive", headers=["Regional"])
    response = client.put(URL, json=_payload(duplicated))
    assert response.status_code == 422 and "Regional" in response.json()["detail"]

    empty = _with(columns, "regional", headers=[" "])
    assert client.put(URL, json=_payload(empty)).status_code == 422

    unknown = [*columns, {**columns[0], "field": "nao_existe"}]
    assert client.put(URL, json=_payload(unknown)).status_code == 422

    # Nome e agência são sempre lidos e obrigatórios.
    locked = _with(columns, "full_name", active=False, required=False)
    saved = client.put(URL, json=_payload(locked)).json()
    name = next(c for c in saved["columns"] if c["field"] == "full_name")
    assert name["active"] and name["required"]


def test_reset_and_sample(client, make_user):
    _admin(client, make_user)
    columns = client.get(URL).json()["columns"]
    client.put(URL, json=_payload(_with(columns, "regional", headers=["Região"])))
    sample = client.get(f"{URL}/sample")
    assert sample.status_code == 200
    header = [c.value for c in load_workbook(io.BytesIO(sample.content)).active[1]]
    assert "Região" in header and "Regional" not in header

    reset = client.delete(URL).json()
    assert reset["is_default"]


def test_only_admin_manages_template(client, make_user):
    make_user("op@example.com", UserRole.OPERATIONAL)
    login(client, "op@example.com")
    assert client.get(URL).status_code == 403
    assert client.put(URL, json={"columns": []}).status_code == 403
